"""Meaningful arithmetic, mapping, source and native-artifact checks.

These checks do not claim CUDA compilation, GPU execution or measured speedup.
"""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json
import math
import re
from openpyxl import load_workbook
from io import BytesIO
from pypdf import PdfReader

BASE = Path(__file__).resolve().parents[1]
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}


def mapping(x, y, w, h, degrees):
    theta = math.radians(degrees)
    c, s = math.cos(theta), math.sin(theta)
    cx, cy = (w-1)/2, (h-1)/2
    return (c*(x-cx)-s*(y-cy)+cx, s*(x-cx)+c*(y-cy)+cy)


def main():
    assert (math.ceil(1024/16), math.ceil(1024/16)) == (64, 64)
    assert 64*64*16*16 == 1048576
    assert 1 / 0.20 == 5
    assert round((950-485)/485*100, 2) == 95.88
    assert math.isclose((2*1.2)/(1*2)-1, .20)
    for w, h in [(1,1), (5,5), (8,6), (1000,750)]:
        gx, gy = math.ceil(w/16), math.ceil(h/16)
        assert gx*16 >= w and gy*16 >= h
        for x, y in [(0,0), (w-1,h-1), ((w-1)/2,(h-1)/2)]:
            sx, sy = mapping(x,y,w,h,0)
            assert math.isclose(sx,x,abs_tol=1e-8) and math.isclose(sy,y,abs_tol=1e-8)
            for degrees in [30, 90, 180, -45]:
                sx, sy = mapping(x,y,w,h,degrees)
                rx, ry = mapping(sx,sy,w,h,-degrees)
                assert math.isclose(rx,x,abs_tol=1e-8) and math.isclose(ry,y,abs_tol=1e-8)
    output = []
    for y in range(5):
        row = []
        for x in range(5):
            sx, sy = mapping(x,y,5,5,90)
            ix, iy = math.floor(sx+.5), math.floor(sy+.5)
            row.append(iy*5+ix)
        output.append(row)
    assert output == [[4,9,14,19,24], [3,8,13,18,23], [2,7,12,17,22], [1,6,11,16,21], [0,5,10,15,20]]

    manifest = json.loads((BASE/'.build/deck_manifest.json').read_text(encoding='utf-8'))
    core_ids = {'transfer', 'rotation', 'worked', 'launch', 'rotation-core',
                'host-code', 'features', 'shape-table', 'speed', 'stream-core',
                'gds', 'hpc', 'environment', 'jobenergy', 'envaccount',
                'governance', 'riskregister', 'fairnessdetail', 'access',
                'accessdetail', 'framework', 'frameworkdetail', 'schedule',
                'conclusions'}
    main_ids = {s['id'] for s in manifest['slides'][:manifest['mainCount']]}
    all_ids = {s['id'] for s in manifest['slides']}
    core_ids.remove('conclusions')
    assert core_ids <= all_ids, core_ids - all_ids
    required_main = {'cover','transfer','rotation','cuda-brief','speed-brief','gds',
                     'environment','governance','access-brief','framework-brief','conclusions'}
    assert main_ids == required_main, main_ids ^ required_main
    main_slides = manifest['slides'][:manifest['mainCount']]
    assert main_slides[-1]['id'] == 'conclusions'
    assert 360 <= sum(s['seconds'] for s in main_slides) <= 420
    assert all(s.get('narration') for s in main_slides)
    # Verify that reorganization retained every existing technical/evidence slide.
    baseline_path = BASE/'.build/pre_timing_manifest.json'
    if baseline_path.exists():
        baseline = json.loads(baseline_path.read_text(encoding='utf-8'))
        current = {s['id']:s for s in manifest['slides']}
        for old in baseline['slides']:
            if old['id'].startswith('references-') or old['id'] == 'ai-declaration':
                continue
            new_id = 'conclusions-detail' if old['id']=='conclusions' else old['id']
            new = current[new_id]
            for key, value in old.items():
                if key not in {'id','label','number'}:
                    assert new[key] == value, (old['id'],key)
    for r in manifest['refs']:
        assert r['url'].startswith('https://'), r['url']
    path = BASE/'Applied2_Combined_Presentation.pptx'
    with ZipFile(path) as z:
        parts = [n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)]
        assert len(parts) == len(manifest['slides'])
        text = '\n'.join(''.join(ET.fromstring(z.read(n)).itertext()) for n in parts)
        for term in ['GPUDirect', 'cuFile', 'SIMT', 'coalescing', 'all-reduce',
                     'dual-use', 'Govern', 'Strubell', 'Gallegos', '35221631', '35512075']:
            assert term.lower() in text.lower(), term
        assert 'provided notebook measures' not in text.lower()
        assert '../../W9/' not in text
        assert 'context/' not in text and 'context\\' not in text
        for n in parts:
            root = ET.fromstring(z.read(n))
            for line in root.findall('.//p:cxnSp/p:spPr/a:ln', NS):
                assert line.find('a:headEnd', NS) is None or line.find('a:headEnd',NS).get('type') == 'none'
                assert line.find('a:tailEnd', NS) is not None
        charts = [n for n in z.namelist() if re.fullmatch(r'.*/chart\d+\.xml',n)]
        values=[]
        for n in charts:
            root=ET.fromstring(z.read(n))
            values.append([float(v.text) for v in root.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v',NS)])
        assert sorted(values) == sorted([[485,950],[.17,.50]]),values
        workbooks=[]
        for n in z.namelist():
            if n.endswith('.xlsx'):
                wb=load_workbook(BytesIO(z.read(n)),data_only=True)
                numeric=[v for row in wb.active.iter_rows(values_only=True) for v in row if isinstance(v,(int,float))]
                workbooks.append(numeric)
        assert any(v==[485,950] for v in workbooks),workbooks
        assert any(v==[.17,.50] for v in workbooks),workbooks
    pdf=BASE/'Applied2_Combined_Presentation.pdf'
    if pdf.exists(): assert len(PdfReader(pdf).pages)==len(parts)
    print('PASS: coordinate inverses, known rotation output, launch arithmetic, chart caches/workbooks, source paths and deck coverage.')
    print('CUDA compilation and GPU runtime remain unverified; no measured performance is asserted.')


if __name__ == '__main__': main()
